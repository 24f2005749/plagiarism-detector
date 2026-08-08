"""Read plain-text, Word, PDF, and scanned/handwritten documents."""

from __future__ import annotations

from pathlib import Path


TEXT_EXTENSIONS = {".txt", ".text", ".md"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS | {".pdf", ".docx"}


class DocumentParseError(ValueError):
    """Raised when a document cannot be converted to text."""


def read_file(path: str | Path) -> str:
    """Extract text from a supported local document."""
    document_path = Path(path)
    if not document_path.is_file():
        raise FileNotFoundError(f"Document not found: {document_path}")

    extension = document_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentParseError(f"Unsupported document type '{extension or '(no extension)'}'. Supported types: {supported}.")
    if extension in TEXT_EXTENSIONS:
        try:
            return document_path.read_text(encoding="utf-8")
        except UnicodeDecodeError as error:
            raise DocumentParseError(f"'{document_path.name}' must be UTF-8 encoded.") from error
    if extension == ".docx":
        return _read_docx(document_path)
    if extension == ".pdf":
        return _read_pdf(document_path)
    return _ocr_image(document_path)


def _read_docx(path: Path) -> str:
    """Extract paragraphs and table cells from a Word document."""
    try:
        from docx import Document
    except ImportError as error:
        raise DocumentParseError("DOCX support requires python-docx. Run: pip install -r requirements.txt") from error
    try:
        document = Document(path)
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        parts.extend(cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip())
    except Exception as error:
        raise DocumentParseError(f"Could not read DOCX '{path.name}': {error}") from error
    return _require_text("\n".join(parts), path.name)


def _read_pdf(path: Path) -> str:
    """Extract a PDF text layer, falling back to OCR for scanned PDFs."""
    try:
        import fitz
    except ImportError as error:
        raise DocumentParseError("PDF support requires PyMuPDF. Run: pip install -r requirements.txt") from error
    try:
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text("text") for page in pdf).strip()
            return _require_text(text, path.name) if has_meaningful_text(text) else _ocr_pdf(pdf, path.name)
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"Could not read PDF '{path.name}': {error}") from error


def _ocr_pdf(pdf: object, filename: str) -> str:
    """Render image-only PDF pages, trying Tesseract before handwriting OCR."""
    try:
        import pytesseract
        from PIL import Image, ImageEnhance, ImageFilter, ImageOps
        rendered_pages = []
        tesseract_pages = []
        for page in pdf:  # type: ignore[union-attr]
            pixmap = page.get_pixmap(dpi=400, alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            rendered_pages.append(image)
            tesseract_pages.append(_preprocess_tesseract_image(image, ImageEnhance, ImageFilter, ImageOps))
        tesseract_text = "\n".join(pytesseract.image_to_string(page, config="--psm 6") for page in tesseract_pages)
        if has_meaningful_text(tesseract_text):
            return _require_text(tesseract_text, filename)
        return _ocr_handwritten_pages(rendered_pages, filename)
    except pytesseract.TesseractNotFoundError as error:
        raise DocumentParseError("Tesseract OCR is not installed. Install it (macOS: brew install tesseract) and retry.") from error
    except ImportError as error:
        raise DocumentParseError("OCR support requires Pillow and pytesseract. Run: pip install -r requirements.txt") from error


def _preprocess_tesseract_image(image: object, image_enhance: object, image_filter: object, image_ops: object) -> object:
    """Keep the existing high-resolution Tesseract preprocessing pipeline."""
    image = image.convert("L")  # type: ignore[union-attr]
    image = image_ops.autocontrast(image)  # type: ignore[union-attr]
    image = image_enhance.Contrast(image).enhance(2)  # type: ignore[union-attr]
    image = image.filter(image_filter.SHARPEN)  # type: ignore[union-attr]
    return image.resize((image.width * 2, image.height * 2))  # type: ignore[union-attr]


def _ocr_handwritten_pages(pages: list[object], filename: str) -> str:
    """Use the lazy TrOCR fallback for pages Tesseract could not read."""
    try:
        from handwriting_ocr import HandwritingOCRError, recognize_handwriting

        text = "\n".join(recognize_handwriting(page) for page in pages)
        return _require_text(text, filename)
    except HandwritingOCRError as error:
        raise DocumentParseError(f"Handwriting OCR failed for '{filename}': {error}") from error


def _ocr_image(path: Path) -> str:
    """OCR a scanned or handwritten image document."""
    try:
        import pytesseract
        from PIL import Image
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image)
        return _require_text(text, path.name)
    except pytesseract.TesseractNotFoundError as error:
        raise DocumentParseError("Tesseract OCR is not installed. Install it (macOS: brew install tesseract) and retry.") from error
    except ImportError as error:
        raise DocumentParseError("OCR support requires Pillow and pytesseract. Run: pip install -r requirements.txt") from error
    except DocumentParseError:
        raise
    except Exception as error:
        raise DocumentParseError(f"Could not OCR image '{path.name}': {error}") from error


def _require_text(text: str, filename: str) -> str:
    text = text.strip()
    if not has_meaningful_text(text):
        raise DocumentParseError(f"No readable text was found in '{filename}'.")
    return text


def has_meaningful_text(text: str) -> bool:
    """Reject empty OCR output and output made entirely of punctuation/symbols."""
    return bool(text.strip()) and any(character.isalnum() for character in text)
